"""
Generates downloadable Excel (.xlsx) and PDF reports for a list of ranked
candidates/resumes - used by both the Bulk Screening history and (via the
same shape) could be reused for job candidate rankings.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

ROW_HEADERS = ["Rank", "Candidate", "Email", "Match %", "Experience (yrs)", "Matched Skills", "Missing Skills"]


def _rows_from_results(results: list[dict]) -> list[list]:
    rows = []
    for i, r in enumerate(results, start=1):
        rows.append([
            i,
            r.get("candidate_name") or r.get("filename") or "Unknown",
            r.get("email") or "-",
            r.get("match_score", 0),
            r.get("experience_years", 0),
            ", ".join(r.get("matched_skills", [])) or "-",
            ", ".join(r.get("missing_critical", [])) or "-",
        ])
    return rows


def build_excel_report(title: str, job_description: str, results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Screening Results"

    ws.merge_cells("A1:G1")
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="4338CA")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = Font(size=9, italic=True, color="666666")

    header_row = 4
    for col, header in enumerate(ROW_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(_rows_from_results(results), start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if c_idx == 4:  # Match %
                score = row[3]
                if score >= 75:
                    cell.fill = PatternFill(start_color="EAF6EE", end_color="EAF6EE", fill_type="solid")
                elif score < 30:
                    cell.fill = PatternFill(start_color="FCEBEB", end_color="FCEBEB", fill_type="solid")

    widths = [6, 22, 24, 10, 16, 40, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def build_shortlist_report_pdf(recruiter_name: str, results: list[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle("Brand", parent=styles["Heading1"], textColor=colors.HexColor("#4338CA"), fontSize=18)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=colors.HexColor("#666666"), fontSize=9)

    elements = [
        Paragraph("Shortlisted Candidates", brand_style),
        Paragraph(
            f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} by {recruiter_name} &middot; Talentum",
            meta_style,
        ),
        Spacer(1, 14),
    ]

    table_data = [["#", "Candidate", "Email", "Job", "Status", "Match %"]]
    for i, r in enumerate(results, start=1):
        table_data.append([
            str(i),
            (r.get("candidate_name") or "Unknown")[:24],
            (r.get("candidate_email") or "-")[:28],
            (r.get("job_title") or "-")[:22],
            (r.get("current_status") or "-").replace("_", " ").title(),
            f"{r['match_score']}%" if r.get("match_score") is not None else "-",
        ])

    table = Table(table_data, colWidths=[1 * cm, 3.5 * cm, 4.5 * cm, 3.5 * cm, 2.5 * cm, 2 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4338CA")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E4E7EC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(table)

    elements.append(Spacer(1, 16))
    elements.append(Paragraph(
        "AI-generated recommendations are decision-support outputs and should not be used as the sole basis "
        "for employment decisions.", meta_style,
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.read()


def build_offer_letter_pdf(
    candidate_name: str,
    job_title: str,
    company_name: str,
    salary: str,
    joining_date: str,
    additional_terms: str | None,
    recruiter_name: str | None,
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2.5 * cm, bottomMargin=2.5 * cm, leftMargin=2.5 * cm, rightMargin=2.5 * cm)
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle("Brand", parent=styles["Heading1"], textColor=colors.HexColor("#4338CA"), fontSize=16)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10.5, leading=16, spaceAfter=10)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=colors.HexColor("#666666"), fontSize=9)

    today = datetime.utcnow().strftime("%d %B %Y")

    elements = [
        Paragraph(company_name, brand_style),
        Paragraph(f"Date: {today}", meta_style),
        Spacer(1, 18),
        Paragraph(f"Dear {candidate_name},", body_style),
        Paragraph(
            f"We are pleased to offer you the position of <b>{job_title}</b> at {company_name}. "
            f"We were impressed by your background and are excited about the possibility of you joining our team.",
            body_style,
        ),
        Paragraph(f"<b>Position:</b> {job_title}", body_style),
        Paragraph(f"<b>Compensation:</b> {salary}", body_style),
        Paragraph(f"<b>Proposed joining date:</b> {joining_date}", body_style),
    ]

    if additional_terms:
        elements.append(Paragraph(f"<b>Additional terms:</b> {additional_terms}", body_style))

    elements.append(Paragraph(
        "This offer is contingent upon successful completion of any standard background or reference checks. "
        "Please confirm your acceptance of this offer at your earliest convenience.",
        body_style,
    ))
    elements.append(Spacer(1, 24))
    elements.append(Paragraph("We look forward to welcoming you to the team.", body_style))
    elements.append(Spacer(1, 24))
    elements.append(Paragraph("Sincerely,", body_style))
    elements.append(Paragraph(recruiter_name or "Hiring Team", body_style))
    elements.append(Paragraph(company_name, body_style))

    elements.append(Spacer(1, 30))
    elements.append(Paragraph(
        "This offer letter was generated by Talentum and should be reviewed by the appropriate parties "
        "before being treated as a binding employment agreement.", meta_style,
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.read()


def build_pdf_report(title: str, job_description: str, results: list[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle("Brand", parent=styles["Heading1"], textColor=colors.HexColor("#4338CA"), fontSize=18)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=colors.HexColor("#666666"), fontSize=9)
    jd_style = ParagraphStyle("JD", parent=styles["Normal"], fontSize=9, leading=13)

    elements = [
        Paragraph(title, brand_style),
        Paragraph(f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} &middot; Talentum Bulk Screening Report", meta_style),
        Spacer(1, 10),
        Paragraph("Job Requirement", styles["Heading3"]),
        Paragraph(job_description.replace("\n", "<br/>")[:1200], jd_style),
        Spacer(1, 14),
        Paragraph("Ranked Results", styles["Heading3"]),
        Spacer(1, 6),
    ]

    table_data = [["#", "Candidate", "Match %", "Exp (yrs)", "Missing Critical Skills"]]
    for i, r in enumerate(results, start=1):
        table_data.append([
            str(i),
            (r.get("candidate_name") or r.get("filename") or "Unknown")[:28],
            f"{r.get('match_score', 0)}%",
            str(r.get("experience_years", 0)),
            ", ".join(r.get("missing_critical", []))[:40] or "-",
        ])

    table = Table(table_data, colWidths=[1.2 * cm, 4.5 * cm, 2.2 * cm, 2.2 * cm, 6.5 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4338CA")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E4E7EC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(table)

    elements.append(Spacer(1, 16))
    elements.append(Paragraph(
        "AI-generated recommendations are decision-support outputs and should not be used as the sole basis "
        "for employment decisions.", meta_style,
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.read()
